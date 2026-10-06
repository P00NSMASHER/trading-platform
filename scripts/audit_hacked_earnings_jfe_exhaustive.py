from __future__ import annotations

import argparse
import base64
import csv
import hashlib
import json
import math
import re
import subprocess
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

SOURCE_REPOSITORY = "vgreg/hacked_earnings_jfe"
SOURCE_COMMIT = "c23c7d79d067a79d70cf20e31b072d3703497eae"
TEXT_EXTENSIONS = {".md", ".txt", ".py", ".csv", ".json", ".ipynb"}
DATASET_STEMS = {"SampleFirms", "TimeOfFirstTrade"}

VENDOR_TERMS = [
    "TAQ", "ITCH", "CBOE", "OptionMetrics", "Optionmetrics", "RavenPack",
    "Ravenpack", "I/B/E/S", "IBES", "WRDS", "NBBO", "CRSP", "Markit",
    "13-F", "Thomson Reuters", "Refinitiv", "EDGAR", "SEC complaint",
    "TimeOfFirstTrade", "first trade", "hacked", "hackers",
]

TIME_RE = re.compile(
    r"(?<!\d)(?:(?:[01]?\d|2[0-3]):[0-5]\d(?::[0-5]\d)?"
    r"(?:\s*(?:a\.?m\.?|p\.?m\.?|ET|EST|EDT|CT|CST|CDT|MT|MST|MDT|PT|PST|PDT))?"
    r"|(?:1[0-2]|0?[1-9])(?:\s*:\s*[0-5]\d)?\s*(?:a\.?m\.?|p\.?m\.?))(?!\w)",
    re.IGNORECASE,
)
ACCEPTANCE_RE = re.compile(
    r"(?:ACCEPTANCE[-_ ]?DATETIME|ACCEPTANCE-DATETIME)\D{0,24}(\d{14})",
    re.IGNORECASE,
)
URL_RE = re.compile(r"https?://[^\s<>\")\]]+", re.IGNORECASE)
FILE_REF_RE = re.compile(
    r"(?:(?:[A-Za-z]:[\\/])|(?:\./|\.\./|/))?"
    r"[A-Za-z0-9_ .()\-\\/]+?\."
    r"(?:h5|hdf|hdf5|csv|parquet|dta|xlsx|xls|pickle|pkl|sas7bdat|zip|txt|json|feather|sav)",
    re.IGNORECASE,
)


def run(*args: str, cwd: Path | None = None, check: bool = True) -> str:
    proc = subprocess.run(
        list(args),
        cwd=str(cwd) if cwd else None,
        check=check,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    return proc.stdout


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def git_blob_sha_bytes(data: bytes) -> str:
    header = f"blob {len(data)}\0".encode("ascii")
    return hashlib.sha1(header + data).hexdigest()


def decode_every_byte(data: bytes) -> tuple[str, str]:
    for enc in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return data.decode(enc), enc
        except UnicodeDecodeError:
            pass
    return data.decode("latin-1"), "latin-1"


def snippet(text: str, start: int, end: int, radius: int = 120) -> str:
    lo = max(0, start - radius)
    hi = min(len(text), end + radius)
    return re.sub(r"\s+", " ", text[lo:hi]).strip()


def scan_text(text: str) -> dict[str, Any]:
    times = list(TIME_RE.finditer(text))
    acceptance = ACCEPTANCE_RE.findall(text)
    lowered = text.lower()
    vendor_counts = {}
    for term in VENDOR_TERMS:
        count = lowered.count(term.lower())
        if count:
            vendor_counts[term] = count
    return {
        "characters": len(text),
        "lines": 0 if not text else text.count("\n") + 1,
        "decoded_utf8_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(),
        "time_token_count": len(times),
        "time_tokens": [m.group(0) for m in times[:200]],
        "time_contexts": [snippet(text, m.start(), m.end()) for m in times[:50]],
        "acceptance_datetime_count": len(acceptance),
        "acceptance_datetimes": acceptance[:200],
        "url_count": len(URL_RE.findall(text)),
        "vendor_term_counts": vendor_counts,
        "file_refs": sorted(set(m.group(0).strip() for m in FILE_REF_RE.finditer(text))),
    }


def scan_release_text(text: str) -> dict[str, Any]:
    times = list(TIME_RE.finditer(text))
    acceptance = ACCEPTANCE_RE.findall(text)
    lowered = text.lower()
    vendor_counts = {}
    for term in VENDOR_TERMS:
        count = lowered.count(term.lower())
        if count:
            vendor_counts[term] = count
    urls = URL_RE.findall(text)
    return {
        "characters": len(text),
        "lines": 0 if not text else text.count("\n") + 1,
        "decoded_utf8_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(),
        "time_token_count": len(times),
        "time_tokens": [m.group(0) for m in times[:200]],
        "time_contexts": [snippet(text, m.start(), m.end()) for m in times[:50]],
        "acceptance_datetime_count": len(acceptance),
        "acceptance_datetimes": acceptance[:200],
        "url_count": len(urls),
        "vendor_term_counts": vendor_counts,
        "file_refs": [],
    }


def tracked_files(source_root: Path) -> list[dict[str, Any]]:
    raw = subprocess.check_output(["git", "ls-files", "-s", "-z"], cwd=source_root)
    rows = []
    for entry in raw.decode("utf-8").split("\0"):
        if not entry:
            continue
        left, path = entry.split("\t", 1)
        mode, blob_sha, stage = left.split()
        data = (source_root / path).read_bytes()
        actual_blob_sha = git_blob_sha_bytes(data)
        if actual_blob_sha != blob_sha:
            raise RuntimeError(
                f"Git blob mismatch for {path}: {actual_blob_sha} != {blob_sha}"
            )
        rows.append(
            {
                "path": path,
                "mode": mode,
                "stage": int(stage),
                "git_blob_sha": blob_sha,
                "size_bytes": len(data),
                "sha256": sha256_bytes(data),
            }
        )
    return rows


def history_inventory(source_root: Path) -> dict[str, Any]:
    head = run("git", "rev-parse", "HEAD", cwd=source_root).strip()
    commits = [x for x in run("git", "rev-list", "--all", cwd=source_root).splitlines() if x]
    refs = [x for x in run(
        "git", "for-each-ref", "--format=%(refname) %(objectname)", cwd=source_root
    ).splitlines() if x]
    tags = [x for x in run("git", "tag", "--list", cwd=source_root).splitlines() if x]
    branches = [x for x in run(
        "git", "branch", "-a", "--format=%(refname:short)", cwd=source_root
    ).splitlines() if x]

    object_rows = []
    historical_paths = set()
    unique_blob_shas = set()
    unique_blob_bytes = 0
    historical_text_blob_scans = []

    for line in run("git", "rev-list", "--objects", "--all", cwd=source_root).splitlines():
        if not line:
            continue
        parts = line.split(" ", 1)
        sha = parts[0]
        path = parts[1] if len(parts) > 1 else ""
        typ = run("git", "cat-file", "-t", sha, cwd=source_root).strip()
        size = int(run("git", "cat-file", "-s", sha, cwd=source_root).strip())
        raw_object = subprocess.check_output(["git", "cat-file", typ, sha], cwd=source_root)
        if len(raw_object) != size:
            raise RuntimeError(f"Git object short-read {sha}: {len(raw_object)} != {size}")
        object_hash = hashlib.sha1(
            f"{typ} {len(raw_object)}\0".encode("ascii") + raw_object
        ).hexdigest()
        if object_hash != sha:
            raise RuntimeError(f"Git object hash mismatch {sha}: got {object_hash}")

        semantic_scan = None
        if typ in {"commit", "tag"}:
            text_value, encoding = decode_every_byte(raw_object)
            semantic_scan = {"encoding": encoding, **scan_text(text_value)}

        if typ == "blob":
            unique_blob_shas.add(sha)
            unique_blob_bytes += size
            if path:
                historical_paths.add(path)
            suffix = Path(path).suffix.lower()
            if path in {".gitignore", "LICENSE"} or suffix in TEXT_EXTENSIONS:
                text_value, encoding = decode_every_byte(raw_object)
                historical_text_blob_scans.append(
                    {
                        "sha": sha,
                        "path": path,
                        "size_bytes": size,
                        "encoding": encoding,
                        **scan_text(text_value),
                    }
                )
        object_rows.append(
            {
                "sha": sha,
                "type": typ,
                "size": size,
                "path": path,
                "raw_sha256": sha256_bytes(raw_object),
                "semantic_scan": semantic_scan,
            }
        )

    current_paths = set(run("git", "ls-files", cwd=source_root).splitlines())
    return {
        "head": head,
        "commit_count": len(commits),
        "commits": commits,
        "refs": refs,
        "tags": tags,
        "branches": branches,
        "reachable_object_count": len(object_rows),
        "reachable_object_bytes_read": sum(row["size"] for row in object_rows),
        "all_reachable_object_hashes_verified": True,
        "unique_blob_count": len(unique_blob_shas),
        "unique_blob_bytes_read": unique_blob_bytes,
        "historical_paths": sorted(historical_paths),
        "historical_only_paths": sorted(historical_paths - current_paths),
        "historical_text_blob_scans": historical_text_blob_scans,
        "objects": object_rows,
    }


def normalize_scalar(value: Any, column: str = "") -> Any:
    try:
        import pandas as pd
        if pd.isna(value):
            return None
    except Exception:
        pass
    if isinstance(value, datetime):
        if column.lower() == "date":
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bytes):
        return value.hex()
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if math.isnan(value):
            return None
        if value.is_integer() and column.lower() in {"permno", "gvkey", "hacked", "actual"}:
            return int(value)
        return float(value)
    s = str(value)
    if column.lower() == "date":
        try:
            import pandas as pd
            return pd.to_datetime(s).date().isoformat()
        except Exception:
            return s
    if "timeoffirsttrade" in column.lower():
        try:
            import pandas as pd
            return pd.to_datetime(s).isoformat(sep=" ")
        except Exception:
            return s
    if column.lower() in {"permno", "gvkey", "hacked", "actual"}:
        try:
            f = float(s)
            if f.is_integer():
                return int(f)
        except Exception:
            pass
    return s


def dataframe_digest(frame: Any) -> tuple[str, list[str]]:
    h = hashlib.sha256()
    columns = [str(c) for c in frame.columns]
    h.update(json.dumps(columns, separators=(",", ":")).encode("utf-8"))
    for row in frame.itertuples(index=False, name=None):
        normalized = [normalize_scalar(value, columns[i]) for i, value in enumerate(row)]
        h.update(json.dumps(
            normalized, ensure_ascii=False, separators=(",", ":")
        ).encode("utf-8"))
        h.update(b"\n")
    return h.hexdigest(), columns


def load_dataset(path: Path) -> Any:
    import pandas as pd
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return pd.read_csv(path)
    if suffix == ".parquet":
        return pd.read_parquet(path)
    if suffix == ".dta":
        return pd.read_stata(path, convert_categoricals=False)
    if suffix == ".xlsx":
        return pd.read_excel(path)
    raise ValueError(path)


def compare_frames(reference: Any, other: Any) -> dict[str, Any]:
    result = {
        "columns_equal": [str(c) for c in reference.columns] == [str(c) for c in other.columns],
        "row_count_equal": len(reference) == len(other),
        "reference_rows": len(reference),
        "other_rows": len(other),
        "mismatch_cells": 0,
        "mismatch_examples": [],
    }
    if not result["columns_equal"] or not result["row_count_equal"]:
        return result

    cols = [str(c) for c in reference.columns]
    for r_idx, (a_row, b_row) in enumerate(zip(
        reference.itertuples(index=False, name=None),
        other.itertuples(index=False, name=None),
    )):
        for c_idx, (a, b) in enumerate(zip(a_row, b_row)):
            col = cols[c_idx]
            na = normalize_scalar(a, col)
            nb = normalize_scalar(b, col)
            if na is None and nb is None:
                equal = True
            elif isinstance(na, (int, float)) and isinstance(nb, (int, float)):
                equal = math.isclose(float(na), float(nb), rel_tol=1e-12, abs_tol=1e-12)
            else:
                equal = na == nb
            if not equal:
                result["mismatch_cells"] += 1
                if len(result["mismatch_examples"]) < 20:
                    result["mismatch_examples"].append(
                        {"row": r_idx + 2, "column": col, "reference": na, "other": nb}
                    )
    return result


def inspect_datasets(source_root: Path) -> dict[str, Any]:
    output = {}
    for stem in sorted(DATASET_STEMS):
        formats = {}
        frames = {}
        for ext in (".csv", ".parquet", ".dta", ".xlsx"):
            path = source_root / "Data" / f"{stem}{ext}"
            frame = load_dataset(path)
            digest, columns = dataframe_digest(frame)
            formats[ext] = {"rows": len(frame), "columns": columns, "canonical_row_digest": digest}
            frames[ext] = frame
        ref = frames[".csv"]
        output[stem] = {
            "formats": formats,
            "csv_comparisons": {
                ext: compare_frames(ref, frames[ext])
                for ext in (".parquet", ".dta", ".xlsx")
            },
        }
    return output


def inspect_xlsx(path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    members = []
    with zipfile.ZipFile(path) as zf:
        bad = zf.testzip()
        if bad is not None:
            raise RuntimeError(f"Corrupt XLSX member {bad} in {path}")
        for info in zf.infolist():
            if info.is_dir():
                continue
            data = zf.read(info)
            row = {
                "name": info.filename,
                "size": len(data),
                "crc": info.CRC,
                "sha256": sha256_bytes(data),
            }
            if info.filename.lower().endswith((".xml", ".rels")):
                text, enc = decode_every_byte(data)
                row["encoding"] = enc
                row.update(scan_text(text))
            members.append(row)

    wb = load_workbook(path, data_only=False, read_only=False, keep_links=True)
    sheets = []
    formula_count = 0
    comment_count = 0
    hyperlink_count = 0
    for ws in wb.worksheets:
        nonempty = 0
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    nonempty += 1
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        formula_count += 1
                    if cell.comment is not None:
                        comment_count += 1
                    if cell.hyperlink is not None:
                        hyperlink_count += 1
        sheets.append(
            {
                "title": ws.title,
                "state": ws.sheet_state,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "nonempty_cells": nonempty,
            }
        )
    return {
        "zip_member_count": len(members),
        "zip_members": members,
        "sheets": sheets,
        "formula_count": formula_count,
        "comment_count": comment_count,
        "hyperlink_count": hyperlink_count,
        "defined_names": [str(item) for item in wb.defined_names.values()],
        "external_link_count": len(getattr(wb, "_external_links", [])),
    }


def inspect_stata(path: Path) -> dict[str, Any]:
    import pandas as pd

    with pd.io.stata.StataReader(path, convert_categoricals=False) as reader:
        variable_labels = reader.variable_labels()
        value_labels = reader.value_labels()
        data_label = getattr(reader, "data_label", None)
        time_stamp = getattr(reader, "time_stamp", None)
        fmtlist = list(getattr(reader, "fmtlist", []) or [])
        varlist = list(getattr(reader, "varlist", []) or [])
        typlist = [str(x) for x in (getattr(reader, "typlist", []) or [])]
    return {
        "data_label": data_label,
        "time_stamp": str(time_stamp) if time_stamp is not None else None,
        "variable_labels": variable_labels,
        "value_labels": value_labels,
        "formats": fmtlist,
        "variables": varlist,
        "types": typlist,
        "metadata_text_scan": scan_text(
            json.dumps(
                {
                    "data_label": data_label,
                    "time_stamp": str(time_stamp) if time_stamp is not None else None,
                    "variable_labels": variable_labels,
                    "value_labels": value_labels,
                    "formats": fmtlist,
                    "variables": varlist,
                    "types": typlist,
                },
                ensure_ascii=False,
                sort_keys=True,
            )
        ),
    }


def inspect_parquet(path: Path) -> dict[str, Any]:
    import pyarrow.parquet as pq

    pf = pq.ParquetFile(path)
    table = pf.read()
    h = hashlib.sha256()
    for batch in table.to_batches(max_chunksize=8192):
        for col_name, column in zip(batch.schema.names, batch.columns):
            h.update(col_name.encode("utf-8"))
            h.update(b"\0")
            for value in column.to_pylist():
                normalized = normalize_scalar(value, col_name)
                h.update(json.dumps(
                    normalized, ensure_ascii=False, separators=(",", ":")
                ).encode("utf-8"))
                h.update(b"\n")
    metadata = {}
    if pf.metadata.metadata:
        for k, v in pf.metadata.metadata.items():
            ks, _ = decode_every_byte(k)
            vs, _ = decode_every_byte(v)
            metadata[ks] = vs
    return {
        "rows": table.num_rows,
        "columns": table.num_columns,
        "schema": str(table.schema),
        "row_groups": pf.metadata.num_row_groups,
        "created_by": pf.metadata.created_by,
        "key_value_metadata": metadata,
        "all_cell_digest": h.hexdigest(),
    }


def walk_strings(obj: Any):
    if isinstance(obj, str):
        yield obj
    elif isinstance(obj, dict):
        for k, v in obj.items():
            yield from walk_strings(k)
            yield from walk_strings(v)
    elif isinstance(obj, list):
        for value in obj:
            yield from walk_strings(value)


def inspect_notebook(path: Path, binary_output_dir: Path | None = None) -> dict[str, Any]:
    raw = path.read_bytes()
    text, enc = decode_every_byte(raw)
    doc = json.loads(text)
    strings = list(walk_strings(doc))
    joined = "\n".join(strings)
    cells = doc.get("cells", [])
    code_cells = [c for c in cells if c.get("cell_type") == "code"]
    markdown_cells = [c for c in cells if c.get("cell_type") == "markdown"]
    output_objects = sum(len(c.get("outputs", [])) for c in code_cells)
    output_texts = []
    embedded_binary_outputs = []
    for cell_index, cell in enumerate(code_cells):
        for output_index, out in enumerate(cell.get("outputs", [])):
            output_texts.extend(walk_strings(out))
            data_payload = out.get("data", {}) if isinstance(out, dict) else {}
            for mime in ("image/png", "image/jpeg", "application/pdf"):
                encoded = data_payload.get(mime)
                if encoded is None:
                    continue
                if isinstance(encoded, list):
                    encoded = "".join(encoded)
                raw_payload = base64.b64decode(encoded, validate=False)
                saved_path = None
                if binary_output_dir is not None:
                    binary_output_dir.mkdir(parents=True, exist_ok=True)
                    extension = {
                        "image/png": ".png",
                        "image/jpeg": ".jpg",
                        "application/pdf": ".pdf",
                    }[mime]
                    safe_stem = re.sub(r"[^A-Za-z0-9_.-]+", "_", path.stem)
                    saved = binary_output_dir / (
                        f"{safe_stem}_cell{cell_index:03d}_output{output_index:02d}{extension}"
                    )
                    saved.write_bytes(raw_payload)
                    saved_path = saved.name
                embedded_binary_outputs.append(
                    {
                        "cell_index": cell_index,
                        "output_index": output_index,
                        "mime": mime,
                        "encoded_characters": len(encoded),
                        "decoded_bytes": len(raw_payload),
                        "sha256": sha256_bytes(raw_payload),
                        "saved_path": saved_path,
                    }
                )
    return {
        "encoding": enc,
        "raw_characters": len(text),
        "json_string_count": len(strings),
        "json_string_characters": sum(len(s) for s in strings),
        "cell_count": len(cells),
        "code_cell_count": len(code_cells),
        "markdown_cell_count": len(markdown_cells),
        "output_object_count": output_objects,
        "output_text_characters": sum(len(s) for s in output_texts),
        "embedded_binary_output_count": len(embedded_binary_outputs),
        "embedded_binary_outputs": embedded_binary_outputs,
        "embedded_binary_output_bytes": sum(
            item["decoded_bytes"] for item in embedded_binary_outputs
        ),
        "joined_string_scan": scan_text(joined),
    }


def inspect_source_files(
    source_root: Path, output_dir: Path
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    rows = tracked_files(source_root)
    special = {}
    for row in rows:
        path = source_root / row["path"]
        suffix = path.suffix.lower()
        if row["path"] in {".gitignore", "LICENSE"} or suffix in TEXT_EXTENSIONS:
            data = path.read_bytes()
            text, enc = decode_every_byte(data)
            row["text_encoding"] = enc
            row.update(scan_text(text))
        if suffix == ".ipynb":
            special[row["path"]] = {
                "notebook": inspect_notebook(
                    path, output_dir / "notebook_embedded_outputs"
                )
            }
        elif suffix == ".xlsx":
            special[row["path"]] = {"xlsx": inspect_xlsx(path)}
        elif suffix == ".parquet":
            special[row["path"]] = {"parquet": inspect_parquet(path)}
        elif suffix == ".dta":
            special[row["path"]] = {"stata": inspect_stata(path)}
    return rows, special


def load_event_map(event_index: Path) -> tuple[dict[str, list[str]], dict[str, list[str]]]:
    event_to_members = {}
    member_to_events = defaultdict(list)
    with event_index.open("r", encoding="utf-8", newline="") as f:
        for row in csv.DictReader(f):
            members = [m for m in row.get("release_member_paths", "").split(";") if m]
            event_to_members[row["event_id"]] = members
            for member in members:
                member_to_events[member].append(row["event_id"])
    return event_to_members, dict(member_to_events)


def inspect_press_archives(source_root: Path, event_index: Path, output_dir: Path) -> dict[str, Any]:
    event_to_members, member_to_events = load_event_map(event_index)
    press_dir = source_root / "Data" / "Press releases"

    member_rows = []
    event_rows = []
    archive_summaries = []
    encoding_counts = Counter()
    vendor_counts = Counter()
    total_members = 0
    total_bytes = 0
    total_chars = 0
    total_lines = 0
    total_times = 0
    total_acceptance = 0
    matched_member_count = 0
    matched_event_times = 0
    matched_event_acceptance = 0
    all_text_digest = hashlib.sha256()
    all_file_refs = set()

    for archive_path in sorted(press_dir.glob("*.zip")):
        with zipfile.ZipFile(archive_path) as zf:
            bad = zf.testzip()
            if bad is not None:
                raise RuntimeError(f"{archive_path.name}: corrupt member {bad}")

            archive_count = 0
            archive_bytes = 0
            archive_chars = 0

            for info in zf.infolist():
                if info.is_dir():
                    continue
                data = zf.read(info)
                text, enc = decode_every_byte(data)
                scan = scan_release_text(text)
                events = member_to_events.get(info.filename, [])

                total_members += 1
                total_bytes += len(data)
                total_chars += scan["characters"]
                total_lines += scan["lines"]
                total_times += scan["time_token_count"]
                total_acceptance += scan["acceptance_datetime_count"]
                archive_count += 1
                archive_bytes += len(data)
                archive_chars += scan["characters"]
                encoding_counts[enc] += 1
                all_file_refs.update(scan["file_refs"])
                for term, count in scan["vendor_term_counts"].items():
                    vendor_counts[term] += count

                all_text_digest.update(archive_path.name.encode("utf-8"))
                all_text_digest.update(b"\0")
                all_text_digest.update(info.filename.encode("utf-8", errors="surrogatepass"))
                all_text_digest.update(b"\0")
                all_text_digest.update(text.encode("utf-8"))

                if events:
                    matched_member_count += 1
                    matched_event_times += scan["time_token_count"]
                    matched_event_acceptance += scan["acceptance_datetime_count"]
                    for event_id in events:
                        event_rows.append(
                            {
                                "event_id": event_id,
                                "archive": archive_path.name,
                                "member": info.filename,
                                "raw_bytes": len(data),
                                "encoding": enc,
                                "characters": scan["characters"],
                                "lines": scan["lines"],
                                "raw_sha256": sha256_bytes(data),
                                "text_sha256": scan["decoded_utf8_sha256"],
                                "time_token_count": scan["time_token_count"],
                                "time_tokens": " || ".join(scan["time_tokens"]),
                                "time_contexts": " || ".join(scan["time_contexts"]),
                                "acceptance_datetime_count": scan["acceptance_datetime_count"],
                                "acceptance_datetimes": " || ".join(scan["acceptance_datetimes"]),
                                "vendor_terms": json.dumps(scan["vendor_term_counts"], sort_keys=True),
                            }
                        )

                member_rows.append(
                    {
                        "archive": archive_path.name,
                        "member": info.filename,
                        "compressed_size": info.compress_size,
                        "raw_bytes": len(data),
                        "crc": f"{info.CRC:08x}",
                        "raw_sha256": sha256_bytes(data),
                        "encoding": enc,
                        "characters": scan["characters"],
                        "lines": scan["lines"],
                        "text_sha256": scan["decoded_utf8_sha256"],
                        "time_token_count": scan["time_token_count"],
                        "acceptance_datetime_count": scan["acceptance_datetime_count"],
                        "url_count": scan["url_count"],
                        "vendor_terms": json.dumps(scan["vendor_term_counts"], sort_keys=True),
                        "event_ids": ";".join(events),
                    }
                )

            archive_summaries.append(
                {
                    "archive": archive_path.name,
                    "member_count": archive_count,
                    "member_uncompressed_bytes_read": archive_bytes,
                    "decoded_characters": archive_chars,
                }
            )

    member_manifest = output_dir / "press_release_member_manifest.csv"
    with member_manifest.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(member_rows[0]))
        writer.writeheader()
        writer.writerows(member_rows)

    event_findings = output_dir / "event_release_text_findings.csv"
    fields = list(event_rows[0]) if event_rows else ["event_id", "archive", "member"]
    with event_findings.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows(event_rows)

    return {
        "archive_count": len(archive_summaries),
        "archives": archive_summaries,
        "member_count": total_members,
        "member_uncompressed_bytes_read": total_bytes,
        "decoded_character_count": total_chars,
        "line_count": total_lines,
        "encoding_counts": dict(encoding_counts),
        "full_decoded_corpus_sha256": all_text_digest.hexdigest(),
        "time_token_count": total_times,
        "acceptance_datetime_count": total_acceptance,
        "vendor_term_counts": dict(vendor_counts),
        "file_refs": sorted(all_file_refs)[:5000],
        "event_count": len(event_to_members),
        "events_with_members": sum(bool(v) for v in event_to_members.values()),
        "events_without_members": sum(not bool(v) for v in event_to_members.values()),
        "matched_member_count": matched_member_count,
        "matched_event_time_token_count": matched_event_times,
        "matched_event_acceptance_datetime_count": matched_event_acceptance,
        "member_manifest": member_manifest.name,
        "event_findings": event_findings.name,
    }


def write_json(path: Path, obj: Any) -> None:
    path.write_text(
        json.dumps(obj, indent=2, sort_keys=True, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )


def flatten_vendor_counts(file_rows: list[dict[str, Any]]) -> dict[str, int]:
    c = Counter()
    for row in file_rows:
        for term, count in row.get("vendor_term_counts", {}).items():
            c[term] += count
    return dict(c)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", type=Path, required=True)
    parser.add_argument("--event-index", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()

    source_root = args.source_root.resolve()
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    head = run("git", "rev-parse", "HEAD", cwd=source_root).strip()
    if head != SOURCE_COMMIT:
        raise RuntimeError(f"Source HEAD {head} != pinned {SOURCE_COMMIT}")

    history = history_inventory(source_root)
    file_rows, special = inspect_source_files(source_root, output_dir)
    datasets = inspect_datasets(source_root)
    press = inspect_press_archives(source_root, args.event_index.resolve(), output_dir)

    current_bytes = sum(int(row["size_bytes"]) for row in file_rows)
    tree_digest = hashlib.sha256()
    for row in sorted(file_rows, key=lambda x: x["path"]):
        tree_digest.update(row["path"].encode("utf-8"))
        tree_digest.update(b"\0")
        tree_digest.update(row["sha256"].encode("ascii"))
        tree_digest.update(b"\n")

    result = {
        "schema_version": "1",
        "source_repository": SOURCE_REPOSITORY,
        "source_commit": SOURCE_COMMIT,
        "coverage_contract": {
            "current_tracked_files_all_read_as_raw_bytes": True,
            "current_tracked_file_count": len(file_rows),
            "current_tracked_bytes_read": current_bytes,
            "current_tree_manifest_sha256": tree_digest.hexdigest(),
            "reachable_git_history_all_unique_blobs_read": True,
            "reachable_git_objects_all_bytes_read_and_hash_verified": history[
                "all_reachable_object_hashes_verified"
            ],
            "reachable_historical_text_blobs_semantically_scanned": len(
                history["historical_text_blob_scans"]
            ),
            "historical_only_paths": history["historical_only_paths"],
            "all_press_release_zip_members_explicitly_decompressed_and_read": True,
            "press_release_member_count": press["member_count"],
            "press_release_uncompressed_bytes_explicitly_read": press["member_uncompressed_bytes_read"],
            "press_release_decoded_characters_scanned": press["decoded_character_count"],
            "dataset_binary_formats_decoded": [".dta", ".parquet", ".xlsx"],
            "jupyter_json_and_outputs_parsed": True,
        },
        "history": history,
        "source_files": file_rows,
        "special_format_inspection": special,
        "dataset_equivalence": datasets,
        "press_release_corpus": press,
        "source_text_vendor_term_counts": flatten_vendor_counts(file_rows),
    }

    write_json(output_dir / "exhaustive_audit.json", result)
    write_json(output_dir / "history_objects.json", history["objects"])
    write_json(output_dir / "special_format_inspection.json", special)
    write_json(output_dir / "dataset_equivalence.json", datasets)

    with (output_dir / "source_file_manifest.csv").open(
        "w", encoding="utf-8", newline=""
    ) as f:
        fields = [
            "path", "mode", "stage", "git_blob_sha", "size_bytes", "sha256",
            "text_encoding", "characters", "lines", "decoded_utf8_sha256",
            "time_token_count", "acceptance_datetime_count", "url_count",
            "vendor_term_counts", "file_refs",
        ]
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for row in file_rows:
            out = dict(row)
            out["vendor_term_counts"] = json.dumps(out.get("vendor_term_counts", {}), sort_keys=True)
            out["file_refs"] = json.dumps(out.get("file_refs", []), ensure_ascii=False)
            writer.writerow(out)

    notebook_refs = {}
    for path, payload in special.items():
        nb = payload.get("notebook")
        if nb:
            notebook_refs[path] = nb["joined_string_scan"]["file_refs"]

    comparison_summary = {}
    for stem, payload in datasets.items():
        comparison_summary[stem] = {
            ext: comp["mismatch_cells"]
            for ext, comp in payload["csv_comparisons"].items()
        }

    summary = [
        "# hacked_earnings_jfe exhaustive audit",
        "",
        f"- Pinned source: {SOURCE_REPOSITORY}@{SOURCE_COMMIT}",
        f"- Current tracked files read byte-for-byte: {len(file_rows)}",
        f"- Current tracked bytes read: {current_bytes:,}",
        f"- Visible Git commits: {history['commit_count']}",
        f"- Reachable unique blobs read: {history['unique_blob_count']}",
        f"- Historical-only paths: {len(history['historical_only_paths'])}",
        f"- Press-release archive members explicitly decompressed/read: {press['member_count']:,}",
        f"- Uncompressed archive bytes explicitly read: {press['member_uncompressed_bytes_read']:,}",
        f"- Decoded archive characters scanned: {press['decoded_character_count']:,}",
        f"- Full decoded archive corpus SHA-256: {press['full_decoded_corpus_sha256']}",
        f"- Events with archived release member(s): {press['events_with_members']}/{press['event_count']}",
        f"- Time-like tokens in all archive text: {press['time_token_count']:,}",
        f"- Time-like tokens in event-matched releases: {press['matched_event_time_token_count']:,}",
        f"- SEC acceptance-datetime tokens in event-matched releases: {press['matched_event_acceptance_datetime_count']:,}",
        "",
        "## Multi-format dataset equivalence",
        json.dumps(comparison_summary, indent=2, sort_keys=True),
        "",
        "## Notebook/file references recovered",
        json.dumps(notebook_refs, indent=2, sort_keys=True, ensure_ascii=False),
        "",
        "## Source-text vendor/data-source term counts",
        json.dumps(result["source_text_vendor_term_counts"], indent=2, sort_keys=True),
        "",
        "## Press-release corpus vendor/data-source term counts",
        json.dumps(press["vendor_term_counts"], indent=2, sort_keys=True),
        "",
        "Detailed manifests are attached as the workflow artifact.",
    ]
    (output_dir / "SUMMARY.md").write_text("\n".join(summary) + "\n", encoding="utf-8")
    print("\n".join(summary))


if __name__ == "__main__":
    main()
